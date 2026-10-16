# -*- coding: utf-8 -*-
"""43 md→docx 转换器（投稿版）：处理 标题/段落/粗体/引用块排除/表格跳过。
转换三个文件：Manuscript / Supplementary Note / Cover Letter。"""
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt

SUB = Path("F:/pe_signature/submission_GB")
SRC = Path(r"D:/zcode/z0927/pe_signature/notes")
JOBS = [
    (SRC / "MANUSCRIPT_GB_FINAL.md", SUB / "Manuscript.docx", True),
    (SUB / "Additional_file_1_Supplementary_Note.md", SUB / "Additional_file_1_Supplementary_Note.docx", False),
    (SUB / "CoverLetter.md", SUB / "CoverLetter.docx", False),
]


def add_runs(par, text):
    """处理 **bold** 片段。"""
    for i, seg in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if not seg:
            continue
        run = par.add_run(seg)
        if i % 2 == 1:
            run.bold = True


def convert(md_path, docx_path, manuscript):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    lines = md_path.read_text(encoding="utf-8").splitlines()
    clean = []
    skip_list = False
    for ln in lines:
        s = ln.strip()
        if s.startswith(">"):            # 内部说明引用块
            continue
        list_markers = ("## 投稿文件清单", "| 文件 | 路径", "| Manuscript（",
                        "| Fig.1–", "| Cover Letter |", "| 代码仓库 |")
        if any(s.startswith(m) for m in list_markers):
            skip_list = True
            continue
        if skip_list and s.startswith("|"):
            continue
        if skip_list and s == "":
            skip_list = False
            continue
        if s.startswith("*(以下为全文") or s.startswith("*(全文") or s.startswith("*(R1–R6"):
            continue
        clean.append(ln)
    for ln in clean:
        if ln.startswith(">"):
            continue
        if "FINAL SUBMISSION COPY" in ln or "MANUSCRIPT —" in ln:
            continue
        if ln.startswith("|") and manuscript:
            continue  # 文末清单表不入正文 docx
        if ln.startswith("## "):
            h = doc.add_heading(re.sub(r"\*\*", "", ln[3:]).strip(), level=1)
            for r in h.runs:
                r.font.name = "Times New Roman"
        elif ln.startswith("### "):
            doc.add_heading(re.sub(r"\*\*", "", ln[4:]).strip(), level=2)
        elif ln.startswith("# "):
            doc.add_heading(re.sub(r"\*\*", "", ln[2:]).strip(), level=0)
        elif ln.startswith("---"):
            continue
        elif ln.strip():
            p = doc.add_paragraph()
            add_runs(p, ln)
    doc.save(str(docx_path))
    print("converted:", docx_path.name)


for src, dst, ms in JOBS:
    convert(src, dst, ms)
print("ALL-DOCX-DONE")
